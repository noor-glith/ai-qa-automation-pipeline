#!/usr/bin/env python3
"""
qa_pipeline.py - requirement -> AI test cases -> Excel/CSV -> Jira -> Playwright -> Bugs

Subcommands
  generate    Read a requirement/user story, ask Claude for test cases, save JSON + XLSX + CSV
  push        Create Epic, Story and test-case issues in Jira from the JSON (idempotent)
  playwright  Turn the test cases into Playwright TypeScript tests (Page Object Model)
  bugs        Read a Playwright JSON report and open Jira Bugs for failed tests (deduplicated)

Examples
  python qa_pipeline.py generate requirements/login.md -o out/login.json
  python qa_pipeline.py push out/login.json --dry-run
  python qa_pipeline.py push out/login.json
  python qa_pipeline.py playwright out/login.json --base-url https://staging.example.com -o tests/
  python qa_pipeline.py bugs test-results/results.json --cases out/login.json

Configuration comes from environment variables (or a .env file) - see .env.example.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from pathlib import Path

import requests
from requests.auth import HTTPBasicAuth

try:
    from dotenv import load_dotenv

    load_dotenv()
except ImportError:
    pass


def env(name: str, default: str | None = None, required: bool = True) -> str:
    val = os.getenv(name, default)
    if required and not val:
        sys.exit(f"Missing environment variable: {name} (see .env.example)")
    return val or ""


# ---------------------------------------------------------------------------
# Claude API
# ---------------------------------------------------------------------------
ANTHROPIC_URL = "https://api.anthropic.com/v1/messages"


def ask_claude(system: str, prompt: str, max_tokens: int = 16000) -> str:
    headers = {
        "x-api-key": env("ANTHROPIC_API_KEY"),
        "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    }
    body = {
        "model": env("ANTHROPIC_MODEL", "claude-sonnet-5-5"),
        "max_tokens": max_tokens,
        "system": system,
        "messages": [{"role": "user", "content": prompt}],
    }
    r = requests.post(ANTHROPIC_URL, headers=headers, json=body, timeout=600)
    if r.status_code != 200:
        sys.exit(f"Claude API error {r.status_code}: {r.text}")
    data = r.json()
    if data.get("stop_reason") == "max_tokens":
        print("Warning: response hit max_tokens and may be truncated.", file=sys.stderr)
    return "".join(b.get("text", "") for b in data["content"] if b["type"] == "text")


def extract_json(text: str) -> dict:
    text = re.sub(r"```(?:json)?", "", text).strip()
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("No JSON object found in model output")
    return json.loads(text[start : end + 1])


# ---------------------------------------------------------------------------
# 1. generate
# ---------------------------------------------------------------------------
GENERATE_SYSTEM = """You are a senior QA engineer. Given a requirement or user story, design
thorough, non-redundant manual test cases covering positive, negative, boundary/edge,
security and accessibility behaviour where relevant.

Respond with ONLY a JSON object - no prose, no markdown fences - in exactly this shape:
{
  "epic":  {"summary": "short feature-level title", "description": "1-3 sentences"},
  "story": {"summary": "As a <user> I want <goal> so that <benefit>",
            "description": "context for developers",
            "acceptance_criteria": ["Given/When/Then ..."]},
  "test_cases": [
    {"id": "TC-001",
     "title": "short imperative title",
     "type": "positive | negative | edge | security | accessibility",
     "priority": "High | Medium | Low",
     "preconditions": "string",
     "steps": ["step 1", "step 2"],
     "test_data": "string (may be empty)",
     "expected_result": "string"}
  ]
}
IDs must be sequential TC-001, TC-002, ... Do not invent requirements that are not implied."""


def validate_cases(data: dict) -> None:
    for key in ("epic", "story", "test_cases"):
        if key not in data:
            raise ValueError(f"AI output missing '{key}'")
    required = {"id", "title", "steps", "expected_result"}
    for i, tc in enumerate(data["test_cases"]):
        missing = required - tc.keys()
        if missing:
            raise ValueError(f"test_cases[{i}] missing {sorted(missing)}")
        if isinstance(tc["steps"], str):
            tc["steps"] = [s for s in tc["steps"].splitlines() if s.strip()]


def cmd_generate(args) -> None:
    requirement = Path(args.requirement).read_text(encoding="utf-8")
    print(f"Asking Claude for test cases ({len(requirement)} chars of requirement)...")
    raw = ask_claude(GENERATE_SYSTEM, f"Requirement / user story:\n\n{requirement}")
    try:
        data = extract_json(raw)
        validate_cases(data)
    except (ValueError, json.JSONDecodeError) as e:
        bad = Path(args.out).with_suffix(".raw.txt")
        bad.parent.mkdir(parents=True, exist_ok=True)
        bad.write_text(raw, encoding="utf-8")
        sys.exit(f"Could not parse AI output ({e}). Raw output saved to {bad}")
    data["source_requirement"] = str(args.requirement)
    save_all(data, Path(args.out))
    print(f"Generated {len(data['test_cases'])} test cases.")


# ---------------------------------------------------------------------------
# Export: JSON + CSV + XLSX
# ---------------------------------------------------------------------------
COLUMNS = ["ID", "Title", "Type", "Priority", "Preconditions", "Steps",
           "Test Data", "Expected Result", "Jira Key", "Automation Status"]


def rows(data: dict) -> list[list[str]]:
    out = []
    for tc in data["test_cases"]:
        steps = "\n".join(f"{n}. {s}" for n, s in enumerate(tc.get("steps", []), 1))
        out.append([tc.get("id", ""), tc.get("title", ""), tc.get("type", ""),
                    tc.get("priority", ""), tc.get("preconditions", ""), steps,
                    tc.get("test_data", ""), tc.get("expected_result", ""),
                    tc.get("jira_key", ""), tc.get("automation_status", "Not automated")])
    return out


def save_all(data: dict, json_path: Path) -> None:
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    csv_path = json_path.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads UTF-8
        w = csv.writer(f)
        w.writerow(COLUMNS)
        w.writerows(rows(data))

    xlsx_path = json_path.with_suffix(".xlsx")
    try:
        write_xlsx(data, xlsx_path)
        print(f"Saved {json_path}, {csv_path}, {xlsx_path}")
    except ImportError:
        print(f"Saved {json_path}, {csv_path} (install openpyxl for .xlsx)")


def write_xlsx(data: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    base = Font(name="Arial", size=10)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill, c.alignment = head, fill, Alignment(vertical="center")
    for r in rows(data):
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = base, wrap

    widths = [9, 40, 13, 10, 30, 60, 25, 45, 12, 18]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Story")
    story = data.get("story", {})
    epic = data.get("epic", {})
    for label, value in [
        ("Epic", epic.get("summary", "")), ("Epic key", epic.get("jira_key", "")),
        ("Story", story.get("summary", "")), ("Story key", story.get("jira_key", "")),
        ("Description", story.get("description", "")),
        ("Acceptance criteria", "\n".join(story.get("acceptance_criteria", []))),
        ("Source requirement", data.get("source_requirement", "")),
    ]:
        info.append([label, value])
    for row in info.iter_rows():
        row[0].font = Font(name="Arial", size=10, bold=True)
        row[1].font, row[1].alignment = base, wrap
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 100
    wb.save(path)


# ---------------------------------------------------------------------------
# Jira client (Cloud, REST API v3)
# ---------------------------------------------------------------------------
def adf_para(text: str) -> dict:
    text = str(text or "").strip()
    return {"type": "paragraph", "content": [{"type": "text", "text": text}] if text else []}


def adf_heading(text: str, level: int = 3) -> dict:
    return {"type": "heading", "attrs": {"level": level},
            "content": [{"type": "text", "text": text}]}


def adf_list(items: list[str], ordered: bool = False) -> dict:
    return {"type": "orderedList" if ordered else "bulletList",
            "content": [{"type": "listItem", "content": [adf_para(i)]}
                        for i in items if str(i).strip()]}


def adf_code(text: str) -> dict:
    return {"type": "codeBlock", "content": [{"type": "text", "text": text or " "}]}


def adf_doc(*nodes) -> dict:
    # Drop empty lists - Jira rejects list nodes with no items
    return {"type": "doc", "version": 1,
            "content": [n for n in nodes if n and n.get("content") != []] or [adf_para(" ")]}


class Jira:
    def __init__(self, dry_run: bool = False):
        self.base = env("JIRA_BASE_URL").rstrip("/")
        self.project = env("JIRA_PROJECT_KEY")
        self.dry_run = dry_run
        self._n = 0
        self.s = requests.Session()
        self.s.auth = HTTPBasicAuth(env("JIRA_EMAIL"), env("JIRA_API_TOKEN"))
        self.s.headers.update({"Accept": "application/json"})

    def _check(self, r: requests.Response) -> dict:
        if not r.ok:
            raise RuntimeError(f"Jira {r.request.method} {r.url} -> {r.status_code}: {r.text}")
        return r.json() if r.text else {}

    def create_issue(self, fields: dict) -> str:
        fields = {"project": {"key": self.project}, **fields}
        if self.dry_run:
            self._n += 1
            key = f"{self.project}-DRY{self._n}"
            print(f"[dry-run] would create {fields['issuetype']['name']}: "
                  f"{fields['summary']} -> {key}")
            return key
        data = self._check(self.s.post(f"{self.base}/rest/api/3/issue", json={"fields": fields}))
        print(f"Created {fields['issuetype']['name']} {data['key']}: {fields['summary']}")
        return data["key"]

    def link(self, from_key: str, to_key: str, link_type: str = "Relates") -> None:
        if self.dry_run:
            print(f"[dry-run] would link {from_key} --{link_type}--> {to_key}")
            return
        self._check(self.s.post(f"{self.base}/rest/api/3/issueLink", json={
            "type": {"name": link_type},
            "inwardIssue": {"key": from_key},
            "outwardIssue": {"key": to_key},
        }))

    def search(self, jql: str, max_results: int = 5) -> list[dict]:
        if self.dry_run:
            return []
        data = self._check(self.s.post(f"{self.base}/rest/api/3/search/jql", json={
            "jql": jql, "maxResults": max_results, "fields": ["summary", "status"]}))
        return data.get("issues", [])

    def comment(self, key: str, body: dict) -> None:
        if self.dry_run:
            print(f"[dry-run] would comment on {key}")
            return
        self._check(self.s.post(f"{self.base}/rest/api/3/issue/{key}/comment",
                                json={"body": body}))

    def attach(self, key: str, path: Path) -> None:
        if self.dry_run or not path.is_file():
            return
        with path.open("rb") as f:
            self._check(self.s.post(f"{self.base}/rest/api/3/issue/{key}/attachments",
                                    headers={"X-Atlassian-Token": "no-check"},
                                    files={"file": (path.name, f)}))


# ---------------------------------------------------------------------------
# 2. push
# ---------------------------------------------------------------------------
def cmd_push(args) -> None:
    path = Path(args.cases)
    data = json.loads(path.read_text(encoding="utf-8"))
    jira = Jira(dry_run=args.dry_run)
    test_type = env("JIRA_TEST_ISSUE_TYPE", "Task")
    labels = ["ai-generated"]

    # Epic
    epic = data["epic"]
    if not epic.get("jira_key"):
        fields = {"summary": epic["summary"], "issuetype": {"name": "Epic"},
                  "description": adf_doc(adf_para(epic.get("description"))), "labels": labels}
        epic_name_field = os.getenv("JIRA_EPIC_NAME_FIELD")  # only old company-managed projects
        if epic_name_field:
            fields[epic_name_field] = epic["summary"]
        epic["jira_key"] = jira.create_issue(fields)

    # Story under the Epic
    story = data["story"]
    if not story.get("jira_key"):
        story["jira_key"] = jira.create_issue({
            "summary": story["summary"][:255],
            "issuetype": {"name": "Story"},
            "parent": {"key": epic["jira_key"]},
            "labels": labels,
            "description": adf_doc(
                adf_para(story.get("description")),
                adf_heading("Acceptance criteria"),
                adf_list(story.get("acceptance_criteria", [])),
            ),
        })

    # Test cases
    for tc in data["test_cases"]:
        if tc.get("jira_key"):
            continue
        fields = {
            "summary": f"[{tc['id']}] {tc['title']}"[:255],
            "issuetype": {"name": test_type},
            "labels": labels + ["test-case",
                                f"priority-{str(tc.get('priority', 'medium')).lower()}",
                                f"type-{str(tc.get('type', 'functional')).lower()}"],
            "description": adf_doc(
                adf_heading("Preconditions"), adf_para(tc.get("preconditions") or "None"),
                adf_heading("Steps"), adf_list(tc.get("steps", []), ordered=True),
                adf_heading("Test data"), adf_para(tc.get("test_data") or "None"),
                adf_heading("Expected result"), adf_para(tc.get("expected_result")),
            ),
        }
        is_subtask = test_type.lower().replace("-", "") == "subtask"
        if is_subtask:
            fields["parent"] = {"key": story["jira_key"]}
        tc["jira_key"] = jira.create_issue(fields)
        if not is_subtask:
            jira.link(tc["jira_key"], story["jira_key"], env("JIRA_TEST_LINK_TYPE", "Relates"))

    if args.dry_run:
        print("Dry run: nothing written to Jira or to the JSON file.")
        return
    save_all(data, path)  # write keys back so re-running skips created issues


# ---------------------------------------------------------------------------
# 3. playwright
# ---------------------------------------------------------------------------
PLAYWRIGHT_SYSTEM = """You are a senior test-automation engineer. Convert manual test cases into
Playwright tests in TypeScript using the Page Object Model.

Rules:
- One page object per page under pages/ (e.g. pages/LoginPage.ts); specs under tests/.
- Prefer user-facing locators: getByRole, getByLabel, getByText, getByTestId. No XPath.
- Use web-first assertions (await expect(locator).toBeVisible()), never fixed waits.
- Every test title MUST start with the test case ID, e.g. test('TC-003 rejects wrong password', ...).
- Read base URL from playwright.config.ts (use page.goto('/path')); read secrets from process.env.
- Where the selectors are unknown, write your best guess and add a // TODO: verify selector comment.
- Include a playwright.config.ts with a json reporter writing test-results/results.json
  alongside the html reporter, and screenshot: 'only-on-failure'.

Output ONLY files, each introduced by a line of the form
=== FILE: relative/path.ts ===
followed by the raw file content. No markdown fences, no commentary."""


def cmd_playwright(args) -> None:
    data = json.loads(Path(args.cases).read_text(encoding="utf-8"))
    cases = [{k: tc.get(k) for k in ("id", "title", "preconditions", "steps",
                                     "test_data", "expected_result")}
             for tc in data["test_cases"]]
    prompt = (f"Base URL: {args.base_url}\n"
              f"Story: {data['story']['summary']}\n\n"
              f"Test cases:\n{json.dumps(cases, indent=2)}")
    raw = ask_claude(PLAYWRIGHT_SYSTEM, prompt)
    out_root = Path(args.out).resolve()
    parts = re.split(r"^=== FILE: (.+?) ===\s*$", raw, flags=re.M)
    if len(parts) < 3:
        sys.exit("Model output had no files. Raw output:\n" + raw[:2000])
    for rel, content in zip(parts[1::2], parts[2::2]):
        target = (out_root / rel.strip()).resolve()
        if out_root not in target.parents:
            print(f"Skipping unsafe path: {rel}")
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        content = re.sub(r"^```\w*\n|```\s*$", "", content.strip()) + "\n"
        target.write_text(content, encoding="utf-8")
        print(f"Wrote {target}")
    print("Review the TODO selectors, then run: npx playwright test")


# ---------------------------------------------------------------------------
# 4. bugs (from Playwright JSON report)
# ---------------------------------------------------------------------------
ANSI = re.compile(r"\x1b\[[0-9;]*m")


def iter_failures(report: dict):
    """Yield failed tests from a Playwright JSON report (reporter: 'json')."""
    def walk(suite, trail):
        trail = trail + [suite.get("title", "")] if suite.get("title") else trail
        for spec in suite.get("specs", []):
            for test in spec.get("tests", []):
                if test.get("status") != "unexpected":  # expected / flaky / skipped are fine
                    continue
                last = (test.get("results") or [{}])[-1]
                err = last.get("error") or (last.get("errors") or [{}])[0]
                yield {
                    "title": spec["title"],
                    "path": " > ".join(t for t in trail if t),
                    "file": f"{spec.get('file', '')}:{spec.get('line', '')}",
                    "project": test.get("projectName", ""),
                    "error": ANSI.sub("", err.get("message", "") or "")[:3000],
                    "stack": ANSI.sub("", err.get("stack", "") or "")[:5000],
                    "attachments": [a.get("path") for a in last.get("attachments", [])
                                    if a.get("path") and a.get("contentType", "").startswith("image/")],
                }
        for child in suite.get("suites", []):
            yield from walk(child, trail)

    for s in report.get("suites", []):
        yield from walk(s, [])


def ci_run_url() -> str:
    if os.getenv("GITHUB_RUN_ID"):
        return (f"{os.getenv('GITHUB_SERVER_URL')}/{os.getenv('GITHUB_REPOSITORY')}"
                f"/actions/runs/{os.getenv('GITHUB_RUN_ID')}")
    return os.getenv("BUILD_URL", "")  # Jenkins


def cmd_bugs(args) -> None:
    report = json.loads(Path(args.report).read_text(encoding="utf-8"))
    jira = Jira(dry_run=args.dry_run)
    case_keys = {}
    if args.cases:
        cases = json.loads(Path(args.cases).read_text(encoding="utf-8"))
        case_keys = {tc["id"]: tc.get("jira_key") for tc in cases["test_cases"]}
    run_url = ci_run_url()
    failures = list(iter_failures(report))
    print(f"{len(failures)} failed test(s) in report.")

    for f in failures:
        # Stable label per test+browser so repeat failures update the same bug
        fp = hashlib.sha1(f"{f['file']}|{f['title']}|{f['project']}".encode()).hexdigest()[:12]
        label = f"pw-{fp}"
        details = adf_doc(
            adf_para(f"Test: {f['path']} > {f['title']}"),
            adf_para(f"Location: {f['file']}   Browser/project: {f['project'] or 'default'}"),
            adf_para(f"CI run: {run_url}") if run_url else None,
            adf_heading("Error"), adf_code(f["error"] or "(no message)"),
            adf_heading("Stack"), adf_code(f["stack"] or "(no stack)"),
        )
        existing = jira.search(f'project = "{jira.project}" AND issuetype = Bug '
                               f'AND labels = "{label}" AND statusCategory != Done', 1)
        if existing:
            key = existing[0]["key"]
            print(f"{key} already open for '{f['title']}' - adding comment")
            jira.comment(key, details)
            continue

        key = jira.create_issue({
            "summary": f"[Automated] {f['title']} failed"
                       f"{' on ' + f['project'] if f['project'] else ''}"[:255],
            "issuetype": {"name": "Bug"},
            "labels": ["automated-test-failure", "playwright", label],
            "description": details,
        })
        for shot in f["attachments"][:3]:
            jira.attach(key, Path(shot))
        m = re.match(r"(TC-\d+)", f["title"])
        if m and case_keys.get(m.group(1)):
            jira.link(key, case_keys[m.group(1)], "Relates")
        elif args.story:
            jira.link(key, args.story, "Relates")

    if failures and args.fail_on_bugs:
        sys.exit(1)


# ---------------------------------------------------------------------------
def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    g = sub.add_parser("generate", help="requirement -> test cases (JSON/XLSX/CSV)")
    g.add_argument("requirement", help="text/markdown file with the requirement or user story")
    g.add_argument("-o", "--out", default="out/test_cases.json")
    g.set_defaults(func=cmd_generate)

    pu = sub.add_parser("push", help="create Epic/Story/test issues in Jira")
    pu.add_argument("cases", help="JSON produced by 'generate'")
    pu.add_argument("--dry-run", action="store_true", help="print what would be created")
    pu.set_defaults(func=cmd_push)

    pw = sub.add_parser("playwright", help="generate Playwright TypeScript POM tests")
    pw.add_argument("cases")
    pw.add_argument("--base-url", required=True)
    pw.add_argument("-o", "--out", default="playwright")
    pw.set_defaults(func=cmd_playwright)

    b = sub.add_parser("bugs", help="create Jira bugs from a Playwright JSON report")
    b.add_argument("report", help="e.g. test-results/results.json")
    b.add_argument("--cases", help="JSON from 'generate'/'push', to link bugs to test cases by TC id")
    b.add_argument("--story", help="fallback issue key to link bugs to, e.g. QA-12")
    b.add_argument("--dry-run", action="store_true")
    b.add_argument("--fail-on-bugs", action="store_true", help="exit 1 if any test failed")
    b.set_defaults(func=cmd_bugs)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
